"""Report outputs and authorized, encrypted supportive supervision evidence."""

import io, base64, hashlib, re
from datetime import date
from xml.sax.saxutils import escape
from fastapi import Depends, HTTPException, UploadFile, File
from fastapi.responses import Response
from sqlalchemy import select
from .models import *
from .schemas import RecordIn
from .security import ADMIN, MANAGERS, AGGREGATE, audit
from .db import get_db
from .config import settings


def install(app):
    from . import main as core

    @app.get("/api/reports/{id}/export/{format}")
    def export_report(id: str, format: str, m=Depends(core.scope), db=Depends(get_db)):
        core.require(m, set(core.ROLES) - {"System Administrator"})
        report = core.owned(db, Report, id, m)
        facility = core.facility(db, report.facility_id, m)
        organization = core.org(db, m)
        registry = {i.id: i for i in db.scalars(core.scoped_query(Indicator, m))}
        rows = []
        for v in report.values:
            i = registry.get(v["indicator_id"])
            if i:
                rows.append(
                    [
                        i.name,
                        i.programme,
                        v["numerator"],
                        v["denominator"],
                        (
                            round(100 * v["numerator"] / v["denominator"], 2)
                            if v["denominator"]
                            else "No estimate"
                        ),
                        i.target,
                        i.approval_reference,
                    ]
                )
        headers = [
            "Indicator",
            "Programme",
            "Numerator",
            "Denominator",
            "Coverage (%)",
            "Target (%)",
            "Approval reference",
        ]
        buf = io.BytesIO()
        if format == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill

            wb = Workbook()
            sheet = wb.active
            sheet.title = "Nutrition report"
            for row in [
                [organization.configuration.get("report_header", organization.name)],
                [facility.name, report.period, report.state, "Revision", report.revision],
                [],
                headers,
            ] + rows:
                sheet.append(row)
            for cell in sheet[4]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="164936")
            for col in ["A", "B", "G"]:
                sheet.column_dimensions[col].width = 40
            for col in ["C", "D", "E", "F"]:
                sheet.column_dimensions[col].width = 18
            # Spreadsheet values from user-supplied text must never be formula cells.
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                        cell.value = "'" + cell.value
            wb.save(buf)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        elif format == "pdf":
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape

            styles = getSampleStyleSheet()
            doc = SimpleDocTemplate(
                buf,
                pagesize=landscape(A4),
                rightMargin=30,
                leftMargin=30,
                topMargin=30,
                bottomMargin=30,
            )
            content = [
                Paragraph("NutriTrack Ghana", styles["Title"]),
                Paragraph(
                    escape(organization.configuration.get("report_header", organization.name)),
                    styles["Heading2"],
                ),
                Paragraph(
                    escape(
                        f"{facility.name} · {report.period} · {report.state} · Revision {report.revision}"
                    ),
                    styles["Normal"],
                ),
                Spacer(1, 18),
            ]
            cells = [
                [Paragraph(escape(str(c)), styles["BodyText"]) for c in r] for r in [headers] + rows
            ]
            table = Table(cells, colWidths=[155, 75, 65, 70, 75, 60, 210], repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2efdf")),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#a7b5ab")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                    ]
                )
            )
            content += [
                table,
                Spacer(1, 16),
                Paragraph(
                    "Estimates use recorded numerators and denominators. No denominator means no estimate. Draft/submitted reports are not official approved results.",
                    styles["Normal"],
                ),
                Paragraph(
                    escape(
                        "Reporting officer: "
                        + organization.configuration.get("reporting_officer", "Not configured")
                    ),
                    styles["Normal"],
                ),
            ]
            doc.build(content)
            content_type = "application/pdf"
        else:
            raise HTTPException(404, "Use pdf or xlsx")
        audit(
            db,
            m,
            "report.exported",
            id,
            {"format": format, "state": report.state, "revision": report.revision},
        )
        db.commit()
        return Response(
            buf.getvalue(),
            media_type=content_type,
            headers={
                "Content-Disposition": f'attachment; filename="nutrition-report-{report.period}.{format}"'
            },
        )

    @app.put("/api/registers/supervision/{id}")
    def supervision_update(id: str, data: RecordIn, m=Depends(core.scope), db=Depends(get_db)):
        core.require(m, ADMIN)
        row = core.owned(db, OperationalRecord, id, m, lock=True)
        if row.kind != "supervision":
            raise HTTPException(404, "Supervision not found")
        core.facility(db, data.facility_id, m)
        try:
            core.validate_record("supervision", data)
        except (ValueError, TypeError):
            raise HTTPException(422, "Invalid supervision values")
        if data.details.get("status") not in {"Scheduled", "Completed", "Cancelled"}:
            raise HTTPException(422, "Invalid visit status")
        allowed_answers = {"Yes", "No", "Not applicable"}
        if any(
            v not in allowed_answers for v in data.details.get("checklist_answers", {}).values()
        ):
            raise HTTPException(422, "Invalid checklist answer")
        audit(db, m, "supervision.updated", id, {"previous_details": row.details})
        row.title = data.title
        row.facility_id = data.facility_id
        row.details = data.details
        row.modified_by = m.user_id
        key = f"supervision:{row.id}"
        if (
            data.details.get("status") == "Completed"
            and data.details.get("corrective_action")
            and not db.scalar(
                select(Action.id).where(
                    Action.organization_id == m.organization_id, Action.signal_key == key
                )
            )
        ):
            db.add(
                Action(
                    **core.provenance(m),
                    title=str(data.details["corrective_action"])[:180],
                    problem=data.details["findings"],
                    facility_id=row.facility_id,
                    due_date=data.details.get("followup_date") or str(date.today()),
                    assigned_to=m.user_id,
                    source=f"Supervision {id}",
                    signal_key=key,
                )
            )
        return core.save(db, row)

    def supervision_record(db, id, m):
        core.require(m, MANAGERS)
        record = core.owned(db, OperationalRecord, id, m)
        if record.kind != "supervision":
            raise HTTPException(404, "Supervision record not found")
        return record

    @app.get("/api/supervision/{id}/evidence")
    def evidence_list(id: str, m=Depends(core.scope), db=Depends(get_db)):
        supervision_record(db, id, m)
        return [
            {k: v for k, v in core.serialize(r).items() if k != "encrypted_content"}
            for r in db.scalars(
                select(Evidence).where(
                    Evidence.organization_id == m.organization_id, Evidence.record_id == id
                )
            )
        ]

    @app.post("/api/supervision/{id}/evidence", status_code=201)
    async def evidence_upload(
        id: str, file: UploadFile = File(...), m=Depends(core.scope), db=Depends(get_db)
    ):
        from cryptography.fernet import Fernet

        supervision_record(db, id, m)
        if not settings().evidence_encryption_key:
            raise HTTPException(503, "Secure evidence storage is not configured")
        content = await file.read(2_000_001)
        if not content or len(content) > 2_000_000:
            raise HTTPException(413, "Provide a file up to 2 MB")
        if content.startswith(b"%PDF-"):
            mime = "application/pdf"
            extension = "pdf"
        elif content.startswith(b"\x89PNG\r\n\x1a\n"):
            mime = "image/png"
            extension = "png"
        elif content.startswith(b"\xff\xd8\xff"):
            mime = "image/jpeg"
            extension = "jpg"
        else:
            raise HTTPException(422, "Only PDF, PNG and JPEG evidence is accepted")
        row = Evidence(
            **core.provenance(m),
            record_id=id,
            filename="evidence." + extension,
            content_type=mime,
            encrypted_content=Fernet(settings().evidence_encryption_key.encode())
            .encrypt(content)
            .decode(),
            size_bytes=len(content),
            sha256=hashlib.sha256(content).hexdigest(),
        )
        db.add(row)
        db.flush()
        audit(
            db, m, "supervision.evidence_uploaded", row.id, {"record_id": id, "bytes": len(content)}
        )
        db.commit()
        return {"id": row.id, "filename": row.filename, "size_bytes": row.size_bytes}

    @app.get("/api/evidence/{id}/download")
    def evidence_download(id: str, m=Depends(core.scope), db=Depends(get_db)):
        from cryptography.fernet import Fernet

        core.require(m, MANAGERS)
        row = core.owned(db, Evidence, id, m)
        supervision_record(db, row.record_id, m)
        if not settings().evidence_encryption_key:
            raise HTTPException(503, "Secure evidence storage is not configured")
        content = Fernet(settings().evidence_encryption_key.encode()).decrypt(
            row.encrypted_content.encode()
        )
        audit(db, m, "supervision.evidence_downloaded", id)
        db.commit()
        return Response(
            content,
            media_type=row.content_type,
            headers={
                "Content-Disposition": f'attachment; filename="{row.filename}"',
                "Content-Security-Policy": "sandbox",
            },
        )
