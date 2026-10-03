"""Bounded import preview for the setup wizard; no records are written."""

import csv, io, json, zipfile
from fastapi import UploadFile, File, Form, Depends, Request, HTTPException
from openpyxl import load_workbook
from .schemas import FacilityIn


def read_rows(content, filename):
    if len(content) > 2_000_000:
        raise HTTPException(413, "Maximum upload is 2 MB")
    if filename.lower().endswith(".csv"):
        return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    if filename.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if len(z.infolist()) > 1000 or sum(i.file_size for i in z.infolist()) > 10_000_000:
                raise HTTPException(413, "Workbook expands beyond the import limit")
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = wb.active
        if sheet.max_row and sheet.max_row > 1001:
            raise HTTPException(422, "Maximum 1000 rows")
        values = list(sheet.values)
        wb.close()
        if not values:
            return []
        return [dict(zip(values[0], r)) for r in values[1:]]
    raise HTTPException(422, "Use CSV or XLSX")


def install(app):
    from . import main as core

    @app.post("/api/setup/facilities/validate")
    async def preview(
        request: Request, file: UploadFile = File(...), subdistricts: str = Form(...)
    ):
        core.limit(request, "setup-import-preview", 5)
        try:
            names = json.loads(subdistricts)
            if (
                not isinstance(names, list)
                or len(names) > 100
                or any(not isinstance(v, str) for v in names)
            ):
                raise ValueError()
            rows = read_rows(await file.read(2_000_001), file.filename or "")
            if not rows or len(rows) > 1000:
                raise HTTPException(422, "Provide 1–1000 rows")
            result = []
            seen = set()
            for number, r in enumerate(rows, 2):
                errors = []
                data = None
                name = str(r.get("facility_name") or "").strip()
                duplicate = name.casefold() in seen
                seen.add(name.casefold())
                if duplicate:
                    errors.append("Duplicate facility name")
                try:
                    sub = str(r.get("subdistrict") or "").strip()
                    if sub not in names:
                        raise ValueError("Unknown sub-district")
                    data = FacilityIn(
                        name=name,
                        code=r.get("facility_code") or None,
                        facility_type=str(r.get("facility_type") or ""),
                        subdistrict_id=sub,
                        community=r.get("community") or None,
                        ownership=r.get("ownership") or "Public",
                        latitude=r.get("latitude") or None,
                        longitude=r.get("longitude") or None,
                        active=str(r.get("status") or "active").lower() == "active",
                    ).model_dump()
                    if data["facility_type"] not in {
                        "Hospital",
                        "Polyclinic",
                        "Health Centre",
                        "CHPS",
                        "Clinic",
                        "Maternity Home",
                        "Other",
                    }:
                        raise ValueError("Choose a configured starting facility type")
                except Exception as e:
                    errors.append(str(e)[:250])
                result.append(
                    dict(
                        id=number,
                        row=number,
                        name=name,
                        errors=errors,
                        status="Duplicate" if duplicate else "Invalid" if errors else "Valid",
                        data=data,
                    )
                )
            return {
                "rows": result,
                "valid": sum(r["status"] == "Valid" for r in result),
                "invalid": sum(r["status"] == "Invalid" for r in result),
                "duplicates": sum(r["status"] == "Duplicate" for r in result),
            }
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(422, "Cannot read file or health sub-district names")
